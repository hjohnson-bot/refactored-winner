#!/usr/bin/env python3
"""
Build MDG_Executive_Dashboard.xlsx — a clean, professional, FORMULA-DRIVEN workbook.

Design:
  * Consolidated: 6 report sheets + the labeled Data tabs that drive them.
  * Dynamic: every number on the report sheets is a live Excel formula
    (SUMIFS / COUNTIFS / structured references) against Excel Tables on the
    Data tabs. Edit the data (or re-run the refresh) and everything recalcs.
  * Sourced: each sheet states where its numbers come from (QuickBooks vs
    Knowify) + the as-of date; a Data Sources panel sits on the dashboard.

Reads the star-schema CSVs in data/ (+ the 2026 invoice pulls in build/raw/).
"""
import csv, os, json, glob
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
RAW  = os.path.join(ROOT, "build", "raw")
OUT  = os.path.join(ROOT, "output", "MDG_Executive_Dashboard.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

def rdrows(name):
    with open(os.path.join(DATA, name), encoding="utf-8") as fh:
        r = list(csv.reader(fh)); return r[0], r[1:]
def rddict(name):
    with open(os.path.join(DATA, name), encoding="utf-8") as fh:
        return list(csv.DictReader(fh))
def n(x):
    try: return float(x)
    except (TypeError, ValueError): return 0.0

recon = json.load(open(os.path.join(DATA, "_reconciliation.json")))
AS_OF = recon["as_of"]
QB, KN = "QuickBooks Online", "Knowify"

# data for python-side layout decisions + status colors (display values are formulas)
gl=rddict("Fact_GL.csv"); wip=rddict("Fact_WIP.csv"); arr=rddict("Fact_AR.csv")
app=rddict("Fact_AP.csv"); cashrow=rddict("Fact_Cash.csv")[0]
bs=rddict("Fact_BalanceSheet.csv"); divs=rddict("Dim_Division.csv")
mhdr, mrows = rdrows("Fact_PL_Monthly.csv")

# invoices (Job, Month, Amount) from the raw 2026 pulls
inv=[]
for p in sorted(glob.glob(os.path.join(RAW,"invoices_2026_p*.json"))):
    for r in json.load(open(p))["Data"]:
        inv.append([r["ProjectName"].strip(), r["InvoiceDate"][:7], round(float(r["TotalAmount"]),2)])

cash_bal=n(cashrow["CashBalance"]); loc_drawn=n(cashrow["LOCDrawn"]); loc_limit=n(cashrow["LOCLimit"])
loc_util=loc_drawn/loc_limit if loc_limit else 0
elig=n(cashrow["EligibleAR"]); adv=n(cashrow["AdvanceRate"]); bbc=min(elig*adv,loc_limit)-loc_drawn
n_fading=sum(1 for r in wip if r["Managed"]=="TRUE" and n(r["ProfitFadePct"])<-0.02)

# ---- clean professional light palette ----
FONT="Segoe UI"
INK="1F2937"; SUBINK="6B7280"; FAINT="9AA3AF"
ACCENT="1F4E79"; ACCENT2="2E75B6"; RULE="D7DCE3"; PANEL="F3F6F9"; CARDBD="C9D2DD"
GOOD="2E7D32"; BAD="C0392B"; WARN="B7791F"; WHITE="FFFFFF"
def fill(c): return PatternFill("solid",fgColor=c)

def put(ws,r,c,v,sz=10,bold=False,color=INK,fmt=None,align="left",fillc=None,italic=False,indent=0):
    cell=ws.cell(row=r,column=c,value=v)
    cell.font=Font(name=FONT,size=sz,bold=bold,color=color,italic=italic)
    cell.alignment=Alignment(horizontal=align,vertical="center",indent=indent,wrap_text=False)
    if fmt: cell.number_format=fmt
    if fillc: cell.fill=fill(fillc)
    return cell

def boxrange(ws,r1,c1,r2,c2,color=CARDBD):
    s=Side(style="thin",color=color)
    for r in range(r1,r2+1):
        for c in range(c1,c2+1):
            ws.cell(row=r,column=c).border=Border(
                left=s if c==c1 else None, right=s if c==c2 else None,
                top=s if r==r1 else None, bottom=s if r==r2 else None)

def title_block(ws,title,sub,source):
    put(ws,1,1,title,18,True,INK,indent=1); ws.merge_cells("A1:N1"); ws.row_dimensions[1].height=30
    put(ws,2,1,sub,10,False,SUBINK,indent=1); ws.merge_cells("A2:N2"); ws.row_dimensions[2].height=16
    for col in range(1,15): ws.cell(row=2,column=col).border=Border(bottom=Side(style="medium",color=ACCENT))
    put(ws,3,1,"Source: "+source,9,False,ACCENT2,italic=True,indent=1); ws.merge_cells("A3:N3"); ws.row_dimensions[3].height=15

def section(ws,row,text,source=None):
    put(ws,row,1,text,12,True,ACCENT); ws.row_dimensions[row].height=20
    if source: put(ws,row,7,"Source: "+source,9,False,FAINT,italic=True)

def thead(ws,row,cols,widths=None,start=1):
    for i,h in enumerate(cols):
        cell=put(ws,row,start+i,h,10,True,WHITE,align=("left" if i==0 else "center"),fillc=ACCENT)
        cell.alignment=Alignment(horizontal=("left" if i==0 else "center"),vertical="center",wrap_text=True,indent=(1 if i==0 else 0))
    ws.row_dimensions[row].height=22
    if widths:
        for i,w in enumerate(widths): ws.column_dimensions[get_column_letter(start+i)].width=w

def trow(ws,row,vals,fmts,bold=False,fillc=None,total=False):
    for i,v in enumerate(vals):
        col=BAD if (isinstance(v,str) and v.startswith("=") and False) else INK
        cell=put(ws,row,1+i,v,10,bold,(WHITE if (fillc==ACCENT) else INK),fmt=fmts[i],
                 align=("left" if i==0 else "right"),fillc=fillc,indent=(1 if i==0 else 0))
    bottom=Side(style="thin",color=(ACCENT if total else RULE))
    for i in range(len(vals)):
        cur=ws.cell(row=row,column=1+i)
        cur.border=Border(bottom=bottom)

wb=Workbook(); wb.remove(wb.active)

# table name refs
GLA='tGL[Amount]';
def sumif_gl(period,cat): return f'SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"{cat}")'
def ni_formula(period):
    return (f'=SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"Revenue")'
            f'-SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"COGS")'
            f'-SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"Operating Expense")'
            f'+SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"Other Income")'
            f'-SUMIFS(tGL[Amount],tGL[Period],"{period}",tGL[Category],"Other Expense")')

# 2026 monthly rows in the Data - Monthly table (header row 3, data from row 4)
m2026_idx=[i for i,r in enumerate(mrows) if r[0][:4]=="2026" and r[0]<="2026-05-31"]
MROW1=4+min(m2026_idx); MROW2=4+max(m2026_idx); MN=len(m2026_idx)
MONTHSHEET="Data - Monthly"

# =====================================================================
# 1) DASHBOARD
# =====================================================================
d=wb.create_sheet("Dashboard"); d.sheet_properties.tabColor=ACCENT; d.sheet_view.showGridLines=False
title_block(d,"Midwest Design Group — Executive Dashboard",
            f"Formula-driven workbook  •  live data  •  as of {AS_OF}","QuickBooks Online + Knowify (see Data tabs)")
# data sources panel
section(d,5,"Data sources")
for r,txt in [(6,f"QuickBooks Online   —   Profit & Loss · Balance Sheet · A/R & A/P Aging · Cash Flow            as of {AS_OF}"),
              (7,f"Knowify   —   Jobs / AJR & Invoices                                                            as of {AS_OF}"),
              (8,"Every figure on the report sheets is a live Excel formula off the Data tabs and reconciles to source (RECONCILIATION.md).")]:
    put(d,r,1,txt,9,False,INK,indent=1,fillc=PANEL); d.merge_cells(start_row=r,start_column=1,end_row=r,end_column=14)
    for c in range(1,15): d.cell(row=r,column=c).fill=fill(PANEL)
boxrange(d,6,1,8,14)

def card(ws,r0,col,label,formula,fmt,context,vcolor=INK,ctxcolor=SUBINK,vsize=18):
    for dr in range(4):
        for dc in (0,1): ws.cell(row=r0+dr,column=col+dc).fill=fill(WHITE)
    put(ws,r0,col,label.upper(),8,True,SUBINK,indent=1); ws.merge_cells(start_row=r0,start_column=col,end_row=r0,end_column=col+1)
    put(ws,r0+1,col,formula,vsize,True,vcolor,fmt=fmt,indent=1); ws.merge_cells(start_row=r0+1,start_column=col,end_row=r0+2,end_column=col+1)
    put(ws,r0+3,col,context,8,False,ctxcolor,indent=1); ws.merge_cells(start_row=r0+3,start_column=col,end_row=r0+3,end_column=col+1)
    boxrange(ws,r0,col,r0+3,col+1)

section(d,10,"Key metrics","live formulas")
rev=f'=SUMIFS(tGL[Amount],tGL[Period],"YTD2026",tGL[Category],"Revenue")'
gm=f'=IFERROR(({sumif_gl("YTD2026","Revenue")}-{sumif_gl("YTD2026","COGS")})/{sumif_gl("YTD2026","Revenue")},0)'
b=11
card(d,b,1,"YTD Revenue",rev,'$#,##0',"QuickBooks · P&L")
card(d,b,3,"Gross Margin",gm,'0.0%',"QuickBooks · P&L")
card(d,b,5,"Net Income",ni_formula("YTD2026"),'$#,##0',"QuickBooks · P&L")
card(d,b,7,"Cash","=SUM(tCash[CashBalance])",'$#,##0',"QuickBooks · balance sheet",vcolor=(BAD if cash_bal<0 else GOOD))
card(d,b,9,"Trade A/R",'=SUMIFS(tAR[Amount],tAR[IsRetainage],"FALSE")','$#,##0',"QuickBooks · A/R aging")
card(d,b,11,"Backlog","=SUM(tWIP[ContractTotal])-SUM(tWIP[Invoiced])",'$#,##0',"Knowify · jobs")
b=16
card(d,b,1,"2026 Forecast Rev",f"=SUM('{MONTHSHEET}'!B{MROW1}:B{MROW2})/{MN}*12",'$#,##0',f"run-rate · {MN} mo complete")
card(d,b,3,"Managed Jobs",'=COUNTIF(tWIP[Managed],"TRUE")','#,##0',"Knowify · PM + budget")
card(d,b,5,"Jobs Fading",'=COUNTIFS(tWIP[Managed],"TRUE",tWIP[ProfitFadePct],"<-0.02")','#,##0',"margin fade < -2%",vcolor=(BAD if n_fading else GOOD))
card(d,b,7,"LOC Utilization","=IFERROR(SUM(tCash[LOCDrawn])/SUM(tCash[LOCLimit]),0)",'0.0%',"QuickBooks + input",vcolor=(BAD if loc_util>=0.9 else (WARN if loc_util>=0.75 else GOOD)))
card(d,b,9,"Overbilled","=SUM(tWIP[Overbilled])",'$#,##0',"Knowify · billed ahead")
card(d,b,11,"BBC Headroom","=MIN(SUM(tCash[EligibleAR])*AVERAGE(tCash[AdvanceRate]),SUM(tCash[LOCLimit]))-SUM(tCash[LOCDrawn])",'$#,##0',"borrowing-base avail",vcolor=(BAD if bbc<0 else GOOD))
for col in range(1,13): d.column_dimensions[get_column_letter(col)].width=11.5

# by-division mini table (formula) + monthly revenue chart
section(d,21,"By division","Knowify · jobs")
thead(d,22,["Division","Contract","Invoiced","Profit $","Margin"],widths=[22,14,14,14,10],start=9)
present_div=[r for r in divs if any(w["DivKey"]==r["DivKey"] for w in wip)]
rr=23
for dv in present_div:
    k=dv["DivKey"]; nm=dv["Division"]
    put(d,rr,9,nm,10,False,INK,indent=1)
    put(d,rr,10,f'=SUMIF(tWIP[DivKey],"{k}",tWIP[ContractTotal])',10,False,INK,fmt='#,##0',align="right")
    put(d,rr,11,f'=SUMIF(tWIP[DivKey],"{k}",tWIP[Invoiced])',10,False,INK,fmt='#,##0',align="right")
    put(d,rr,12,f'=SUMIF(tWIP[DivKey],"{k}",tWIP[ProfitAmount])',10,False,INK,fmt='#,##0',align="right")
    put(d,rr,13,f'=IFERROR(SUMIF(tWIP[DivKey],"{k}",tWIP[ProfitAmount])/SUMIF(tWIP[DivKey],"{k}",tWIP[ContractTotal]),0)',10,False,INK,fmt='0.0%',align="right")
    for c in range(9,14): d.cell(row=rr,column=c).border=Border(bottom=Side(style="thin",color=RULE))
    rr+=1
put(d,22,1,"Revenue by month (2026)",11,True,ACCENT)

# =====================================================================
# 2) P&L
# =====================================================================
pl=wb.create_sheet("P&L"); pl.sheet_properties.tabColor=ACCENT; pl.sheet_view.showGridLines=False
title_block(pl,"Profit & Loss","Company P&L by period — every cell is a SUMIFS over the Data - GL table",
            "QuickBooks Online — Profit & Loss")
PERIODS=[("FY2024","FY 2024"),("FY2025","FY 2025"),("YTD2026",f"YTD 2026")]
thead(pl,5,["Line item"]+[lbl for _,lbl in PERIODS],widths=[40,18,18,18])
def pl_rowvals(maker): return [maker(pk) for pk,_ in PERIODS]
r=6
def pline(label,vals,fmt,bold=False,total=False):
    global r
    put(pl,r,1,label,10,bold,INK,indent=1,fillc=(PANEL if total else None))
    for i,v in enumerate(vals):
        put(pl,r,2+i,v,10,bold,INK,fmt=fmt,align="right",fillc=(PANEL if total else None))
    for c in range(1,5): pl.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=(ACCENT if total else RULE)))
    r+=1
pline("Revenue",pl_rowvals(lambda pk:f'={sumif_gl(pk,"Revenue")}'),'$#,##0',bold=True)
pline("Cost of goods sold",pl_rowvals(lambda pk:f'={sumif_gl(pk,"COGS")}'),'$#,##0')
gp_rows=r
pline("Gross profit",[f'={get_column_letter(2+i)}{r-2}-{get_column_letter(2+i)}{r-1}' for i in range(3)],'$#,##0',bold=True,total=True)
pline("Gross margin %",[f'=IFERROR({get_column_letter(2+i)}{r-1}/{get_column_letter(2+i)}{r-3},0)' for i in range(3)],'0.0%')
pline("Operating expenses",pl_rowvals(lambda pk:f'={sumif_gl(pk,"Operating Expense")}'),'$#,##0')
pline("Other income",pl_rowvals(lambda pk:f'={sumif_gl(pk,"Other Income")}'),'$#,##0')
pline("Other expense",pl_rowvals(lambda pk:f'={sumif_gl(pk,"Other Expense")}'),'$#,##0')
gprow=gp_rows  # row index of Gross profit
# Net income = GP - OpEx + OtherInc - OtherExp  (reference rows)
ni_vals=[f'={get_column_letter(2+i)}{gprow}-{get_column_letter(2+i)}{gprow+2}+{get_column_letter(2+i)}{gprow+3}-{get_column_letter(2+i)}{gprow+4}' for i in range(3)]
pline("Net income",ni_vals,'$#,##0',bold=True,total=True)
pline("Net margin %",[f'=IFERROR({get_column_letter(2+i)}{r-1}/{get_column_letter(2+i)}{gprow-2},0)' for i in range(3)],'0.0%')
put(pl,r+1,1,"FY2024 / FY2025 are closed-year QuickBooks P&L; YTD2026 ties to QuickBooks to the penny (see RECONCILIATION.md).",9,False,FAINT,italic=True)

# =====================================================================
# 3) PROJECTS (Knowify) — formula rollups by PM and division
# =====================================================================
pj=wb.create_sheet("Projects"); pj.sheet_properties.tabColor=ACCENT; pj.sheet_view.showGridLines=False
title_block(pj,"Projects — WIP & Profit Fade","Per-PM and per-division rollups (formulas over Data - WIP); job-level detail on the Data - WIP tab",
            "Knowify — Jobs / AJR")
section(pj,5,"By project manager","Knowify")
thead(pj,6,["Project manager","Jobs","Contract","Invoiced","Profit $","Avg fade %"],widths=[24,8,15,15,15,11])
pms=sorted({w["PMName"] for w in wip}, key=lambda p:-sum(n(w["ContractTotal"]) for w in wip if w["PMName"]==p))
r=7
for pm in pms:
    pmq=pm.replace('"','""')
    put(pj,r,1,pm,10,False,INK,indent=1)
    put(pj,r,2,f'=COUNTIF(tWIP[PMName],"{pmq}")',10,False,INK,fmt='#,##0',align="right")
    put(pj,r,3,f'=SUMIF(tWIP[PMName],"{pmq}",tWIP[ContractTotal])',10,False,INK,fmt='#,##0',align="right")
    put(pj,r,4,f'=SUMIF(tWIP[PMName],"{pmq}",tWIP[Invoiced])',10,False,INK,fmt='#,##0',align="right")
    put(pj,r,5,f'=SUMIF(tWIP[PMName],"{pmq}",tWIP[ProfitAmount])',10,False,INK,fmt='#,##0',align="right")
    put(pj,r,6,f'=IFERROR(AVERAGEIFS(tWIP[ProfitFadePct],tWIP[PMName],"{pmq}",tWIP[Managed],"TRUE"),0)',10,False,INK,fmt='0.0%',align="right")
    for c in range(1,7): pj.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE))
    r+=1
put(pj,r,1,"Total",10,True,INK,indent=1,fillc=PANEL)
for i,(cl,fmt) in enumerate([("B",'#,##0'),("C",'#,##0'),("D",'#,##0'),("E",'#,##0')]):
    put(pj,r,2+i,f'=SUBTOTAL(9,{cl}7:{cl}{r-1})',10,True,INK,fmt=fmt,align="right",fillc=PANEL)
put(pj,r,6,"",10,True,fillc=PANEL)
for c in range(1,7): pj.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=ACCENT))
note=r+2
put(pj,note,1,"Tip: open the Data - WIP tab to sort/filter every job (profit fade, WIP net, % complete, open A/R).",9,False,FAINT,italic=True)

# =====================================================================
# 4) RECEIVABLES & PAYABLES
# =====================================================================
rp=wb.create_sheet("Receivables & Payables"); rp.sheet_properties.tabColor=ACCENT; rp.sheet_view.showGridLines=False
title_block(rp,"Receivables & Payables","Aging by bucket (formulas over Data - AR / Data - AP)","QuickBooks Online — A/R & A/P Aging")
ARB=["Current","1-30","31-60","61-90","91+"]
section(rp,5,"Accounts receivable — aging","QuickBooks")
thead(rp,6,["Bucket","Amount"],widths=[24,18])
r=7
for bkt in ARB:
    put(rp,r,1,bkt,10,False,INK,indent=1)
    put(rp,r,2,f'=SUMIFS(tAR[Amount],tAR[Bucket],"{bkt}",tAR[IsRetainage],"FALSE")',10,False,INK,fmt='$#,##0',align="right")
    for c in (1,2): rp.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE)); r_=r
    r+=1
put(rp,r,1,"Total trade A/R",10,True,INK,indent=1,fillc=PANEL)
put(rp,r,2,'=SUMIFS(tAR[Amount],tAR[IsRetainage],"FALSE")',10,True,INK,fmt='$#,##0',align="right",fillc=PANEL)
for c in (1,2): rp.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=ACCENT))
r+=1
put(rp,r,1,"Retainage receivable (held)",10,False,INK,indent=1)
put(rp,r,2,'=SUMIFS(tAR[Amount],tAR[IsRetainage],"TRUE")',10,False,ACCENT2,fmt='$#,##0',align="right"); r+=1
put(rp,r,1,"A/R as % of YTD revenue",10,False,SUBINK,indent=1)
put(rp,r,2,f'=IFERROR(SUMIFS(tAR[Amount],tAR[IsRetainage],"FALSE")/{sumif_gl("YTD2026","Revenue")},0)',10,False,SUBINK,fmt='0.0%',align="right")
apr=r+3
section(rp,apr,"Accounts payable — aging","QuickBooks")
thead(rp,apr+1,["Bucket","Amount"],widths=[24,18])
r=apr+2
for bkt in ARB:
    put(rp,r,1,bkt,10,False,INK,indent=1)
    put(rp,r,2,f'=SUMIFS(tAP[Amount],tAP[Bucket],"{bkt}")',10,False,INK,fmt='$#,##0',align="right")
    for c in (1,2): rp.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE))
    r+=1
put(rp,r,1,"Total A/P",10,True,INK,indent=1,fillc=PANEL)
put(rp,r,2,'=SUM(tAP[Amount])',10,True,INK,fmt='$#,##0',align="right",fillc=PANEL)
for c in (1,2): rp.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=ACCENT))

# =====================================================================
# 5) CASH & BALANCE SHEET
# =====================================================================
cb=wb.create_sheet("Cash & Balance Sheet"); cb.sheet_properties.tabColor=ACCENT; cb.sheet_view.showGridLines=False
title_block(cb,"Cash & Liquidity / Balance Sheet","Formulas over Data - Cash and Data - Balance Sheet  •  yellow = user-maintained inputs",
            "QuickBooks Online — Balance Sheet & Cash Flow")
section(cb,5,"Cash & liquidity","QuickBooks + inputs")
liq=[("Cash balance","=SUM(tCash[CashBalance])",'$#,##0',None),
     ("Operating cash flow (YTD)","=SUM(tCash[OperatingCF])",'$#,##0',None),
     ("Investing cash flow (YTD)","=SUM(tCash[InvestingCF])",'$#,##0',None),
     ("Financing cash flow (YTD)","=SUM(tCash[FinancingCF])",'$#,##0',None),
     ("LOC drawn","=SUM(tCash[LOCDrawn])",'$#,##0',None),
     ("LOC limit  (input)","=SUM(tCash[LOCLimit])",'$#,##0',"input"),
     ("LOC utilization","=IFERROR(SUM(tCash[LOCDrawn])/SUM(tCash[LOCLimit]),0)",'0.0%',None),
     ("Advance rate  (input)","=AVERAGE(tCash[AdvanceRate])",'0.0%',"input"),
     ("Eligible A/R  (input)","=SUM(tCash[EligibleAR])",'$#,##0',"input"),
     ("BBC availability (headroom)","=MIN(SUM(tCash[EligibleAR])*AVERAGE(tCash[AdvanceRate]),SUM(tCash[LOCLimit]))-SUM(tCash[LOCDrawn])",'$#,##0',None),
     ("Liquidity status",'=IF(IFERROR(SUM(tCash[LOCDrawn])/SUM(tCash[LOCLimit]),0)>=0.9,"Critical",IF(IFERROR(SUM(tCash[LOCDrawn])/SUM(tCash[LOCLimit]),0)>=0.75,"Watch","Healthy"))','General',None)]
r=6
for label,formula,fmt,flag in liq:
    put(cb,r,1,label,10,False,INK,indent=1)
    cell=put(cb,r,2,formula,10,bool(flag),INK,fmt=fmt,align="right",fillc=("FFF3CD" if flag else None))
    for c in (1,2): cb.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE))
    r+=1
bsr=r+2
section(cb,bsr,"Balance sheet","QuickBooks")
thead(cb,bsr+1,["Account","Amount"],widths=[40,18])
r=bsr+2
for secname in ["Assets","Liabilities","Equity"]:
    for row in [x for x in bs if x["Section"]==secname]:
        acct=row["Account"].replace('"','""')
        put(cb,r,1,row["Account"],10,False,INK,indent=2)
        put(cb,r,2,f'=SUMIFS(tBS[Amount],tBS[Account],"{acct}",tBS[Section],"{secname}")',10,False,INK,fmt='$#,##0',align="right")
        for c in (1,2): cb.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE))
        r+=1
    put(cb,r,1,f"Total {secname.lower()}",10,True,INK,indent=1,fillc=PANEL)
    put(cb,r,2,f'=SUMIFS(tBS[Amount],tBS[Section],"{secname}")',10,True,INK,fmt='$#,##0',align="right",fillc=PANEL)
    for c in (1,2): cb.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=ACCENT))
    r+=2

# =====================================================================
# 6) JOB REVENUE BY MONTH (Knowify invoices)
# =====================================================================
jr=wb.create_sheet("Job Revenue by Month"); jr.sheet_properties.tabColor=ACCENT; jr.sheet_view.showGridLines=False
title_block(jr,"Revenue Billed by Job — 2026 (Jan–May)","Each cell is a SUMIFS over the Data - Invoices table (by invoice date)",
            "Knowify — Invoices")
MONTHS=[("2026-01","Jan"),("2026-02","Feb"),("2026-03","Mar"),("2026-04","Apr"),("2026-05","May")]
thead(jr,5,["Job"]+[m for _,m in MONTHS]+["Total Jan–May"],widths=[52,13,13,13,13,13,15])
# job list sorted by total desc
jt={}
for j,m,a in inv: jt[j]=jt.get(j,0)+a
jobs=sorted(jt, key=lambda j:-jt[j])
r=6
for j in jobs:
    jq=j.replace('"','""')
    put(jr,r,1,j,10,False,INK,indent=1)
    for i,(mk,_) in enumerate(MONTHS):
        put(jr,r,2+i,f'=SUMIFS(tInv[Amount],tInv[Job],"{jq}",tInv[Month],"{mk}")',10,False,INK,fmt='#,##0',align="right")
    put(jr,r,7,f'=SUM(B{r}:F{r})',10,True,INK,fmt='#,##0',align="right")
    for c in range(1,8): jr.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=RULE))
    r+=1
put(jr,r,1,"Total — all jobs",10,True,INK,indent=1,fillc=PANEL)
for c in range(2,8):
    L=get_column_letter(c); put(jr,r,c,f'=SUM({L}6:{L}{r-1})',10,True,INK,fmt='#,##0',align="right",fillc=PANEL)
for c in range(1,8): jr.cell(row=r,column=c).border=Border(bottom=Side(style="thin",color=ACCENT))
jr.freeze_panes="B6"

# =====================================================================
# DATA TABS (the engine)
# =====================================================================
def data_tab(sheetname, tablename, csvfile, source, money_cols=(), pct_cols=(), id_cols=()):
    headers, rows = rdrows(csvfile)
    ws=wb.create_sheet(sheetname); ws.sheet_properties.tabColor="9AA3AF"; ws.sheet_view.showGridLines=False
    put(ws,1,1,f"DATA · {tablename}    —    Source: {source}    —    as of {AS_OF}    (drives the report sheets by formula)",9,True,SUBINK)
    hr=3
    for j,h in enumerate(headers,1):
        cell=put(ws,hr,j,h,9,True,WHITE,fillc=ACCENT2,align="center"); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for i,row in enumerate(rows,hr+1):
        for j,v in enumerate(row,1):
            h=headers[j-1]
            val=v if h in id_cols else (n(v) if (h in money_cols or h in pct_cols) else (n(v) if _isnum(v) else v))
            cell=put(ws,i,j,val,9,False,INK,align=("right" if (h in money_cols or h in pct_cols) else "left"))
            if h in money_cols: cell.number_format='#,##0'
            elif h in pct_cols: cell.number_format='0.0%'
    last=get_column_letter(len(headers)); end=hr+len(rows)
    t=Table(displayName=tablename, ref=f"A{hr}:{last}{end}")
    t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True)
    ws.add_table(t)
    for j,h in enumerate(headers,1): ws.column_dimensions[get_column_letter(j)].width=max(10,min(46,len(h)+3))
    ws.freeze_panes=f"A{hr+1}"
    return ws

def _isnum(v):
    try: float(v); return True
    except (TypeError,ValueError): return False

# write invoices CSV-like rows directly
ws_inv=wb.create_sheet("Data - Invoices"); ws_inv.sheet_properties.tabColor="9AA3AF"; ws_inv.sheet_view.showGridLines=False
put(ws_inv,1,1,f"DATA · tInv    —    Source: {KN} — Invoices    —    as of {AS_OF}    (drives Job Revenue by Month)",9,True,SUBINK)
for j,h in enumerate(["Job","Month","Amount"],1):
    cell=put(ws_inv,3,j,h,9,True,WHITE,fillc=ACCENT2,align="center")
for i,(j_,m_,a_) in enumerate(inv,4):
    put(ws_inv,i,1,j_,9,False,INK); put(ws_inv,i,2,m_,9,False,INK)
    put(ws_inv,i,3,a_,9,False,INK,fmt='#,##0',align="right")
tinv=Table(displayName="tInv", ref=f"A3:C{3+len(inv)}"); tinv.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True)
ws_inv.add_table(tinv)
ws_inv.column_dimensions["A"].width=52; ws_inv.column_dimensions["B"].width=10; ws_inv.column_dimensions["C"].width=14; ws_inv.freeze_panes="A4"

GL_MONEY={"Amount"}
data_tab("Data - GL","tGL","Fact_GL.csv",f"{QB} — Profit & Loss",money_cols={"Amount"},id_cols={"Account","Period","PeriodEnd","DivKey","Category"})
wsm=data_tab(MONTHSHEET,"tMonthly","Fact_PL_Monthly.csv",f"{QB} — monthly P&L",
         money_cols={"Revenue","COGS","GrossProfit","OpEx","NetIncome"},pct_cols={"GrossMarginPct","NetMarginPct"},id_cols={"MonthEnd"})
data_tab("Data - WIP","tWIP","Fact_WIP.csv",f"{KN} — Jobs / AJR",
         money_cols={"ContractTotal","ChangeOrders","Invoiced","PaymentsInvoices","BudgetTotal","ActualCost","EarnedRevenue","WIPNet","KnowifyWIP","Overbilled","Underbilled","ProfitAmount","ProjectedProfit","Retainage","OpenAR"},
         pct_cols={"PctComplete","ProfitPct","ProjectedProfitPct","EstMarginPct","ProfitFadePct"},
         id_cols={"SnapshotDate","ProjectId","Job","DivKey","PMName","Customer","Status","HasBudget","Managed"})
data_tab("Data - AR","tAR","Fact_AR.csv",f"{QB} — A/R aging",money_cols={"Amount"},id_cols={"Customer","Bucket","IsRetainage","AsOfDate"})
data_tab("Data - AP","tAP","Fact_AP.csv",f"{QB} — A/P aging",money_cols={"Amount"},id_cols={"Vendor","Bucket","AsOfDate"})
data_tab("Data - Cash","tCash","Fact_Cash.csv",f"{QB} — balance sheet & cash flow",
         money_cols={"CashBalance","LOCDrawn","LOCCapX","OperatingCF","InvestingCF","FinancingCF","NetCashChange","LOCLimit","EligibleAR"},
         pct_cols={"AdvanceRate"},id_cols={"Date"})
data_tab("Data - Balance Sheet","tBS","Fact_BalanceSheet.csv",f"{QB} — balance sheet",money_cols={"Amount"},id_cols={"Account","Section","AsOfDate"})

# Dashboard monthly revenue chart (references the Data - Monthly table)
ch=LineChart(); ch.title="Revenue by month — 2026"; ch.height=7.2; ch.width=15.5; ch.legend=None
ch.add_data(Reference(wsm,min_col=2,min_row=MROW1,max_row=MROW2)); ch.set_categories(Reference(wsm,min_col=1,min_row=MROW1,max_row=MROW2))
ch.y_axis.majorGridlines=None
for s in ch.series:
    s.graphicalProperties=GraphicalProperties(ln=LineProperties(solidFill=ACCENT,w=28000))
d.add_chart(ch,"A23")

wb.save(OUT)
print("Saved", OUT)
print("Sheets:", wb.sheetnames)
